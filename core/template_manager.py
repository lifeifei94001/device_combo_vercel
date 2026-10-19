# -*- coding: utf-8 -*-
"""
模板管理与智能精确对齐导入引擎 (纯 openpyxl + csv 实现，支持文件路径与 BytesIO 内存流，零 pandas 依赖)
"""
import os
import io
import csv
import re
import openpyxl
from typing import List, Dict, Any, Tuple
from config.constants import (
    ANDROID_INPUT_COLUMNS, ANDROID_SAMPLE_DATA,
    LAPTOP_INPUT_COLUMNS, LAPTOP_SAMPLE_DATA,
    WATCH_INPUT_COLUMNS, WATCH_SAMPLE_DATA
)


class TemplateManager:
    @staticmethod
    def generate_template(module_id: str, output_target: Any) -> bool:
        """生成标准 Excel 模版文件 (支持保存到文件路径或 BytesIO 内存流)"""
        try:
            if isinstance(output_target, str):
                os.makedirs(os.path.dirname(output_target), exist_ok=True)

            if module_id == "android":
                cols, sample = ANDROID_INPUT_COLUMNS, ANDROID_SAMPLE_DATA
            elif module_id == "laptop":
                cols, sample = LAPTOP_INPUT_COLUMNS, LAPTOP_SAMPLE_DATA
            elif module_id == "watch":
                cols, sample = WATCH_INPUT_COLUMNS, WATCH_SAMPLE_DATA
            else:
                cols, sample = ANDROID_INPUT_COLUMNS, ANDROID_SAMPLE_DATA

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "机型配置填报模版"

            ws.append(list(cols))
            for item in sample:
                ws.append([item.get(c, '') for c in cols])

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len * 2, 14)

            wb.save(output_target)
            wb.close()
            return True
        except Exception as e:
            print(f"生成模板失败: {e}")
            return False

    @staticmethod
    def _read_rows_from_input(file_path_or_bytes: Any) -> Tuple[List[str], List[Dict[str, str]]]:
        """从 bytes 或文件路径读取表头与数据行字典列表 (支持 .xlsx 与 .csv)"""
        if isinstance(file_path_or_bytes, bytes):
            # 判断是否为 xlsx (PK\x03\x04 魔数)
            if file_path_or_bytes[:4] == b'PK\x03\x04':
                wb = openpyxl.load_workbook(io.BytesIO(file_path_or_bytes), read_only=True, data_only=True)
                ws = wb.active
                raw_rows = list(ws.iter_rows(values_only=True))
                wb.close()
            else:
                # 按 CSV 解析 (自适应 UTF-8-SIG / UTF-8 / GBK)
                text = None
                for enc in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030'):
                    try:
                        text = file_path_or_bytes.decode(enc)
                        break
                    except Exception:
                        continue
                if text is None:
                    text = file_path_or_bytes.decode('utf-8', errors='ignore')
                raw_rows = list(csv.reader(io.StringIO(text)))
        elif isinstance(file_path_or_bytes, str):
            if file_path_or_bytes.lower().endswith('.csv'):
                text = None
                for enc in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030'):
                    try:
                        with open(file_path_or_bytes, 'r', encoding=enc) as f:
                            text = f.read()
                        break
                    except Exception:
                        continue
                raw_rows = list(csv.reader(io.StringIO(text or '')))
            else:
                wb = openpyxl.load_workbook(file_path_or_bytes, read_only=True, data_only=True)
                ws = wb.active
                raw_rows = list(ws.iter_rows(values_only=True))
                wb.close()
        else:
            raise ValueError("未知的文件格式")

        if not raw_rows:
            return [], []

        headers = [str(c if c is not None else '').strip() for c in raw_rows[0]]
        dict_rows = []
        for r in raw_rows[1:]:
            if not r:
                continue
            row_dict = {}
            for idx, col_name in enumerate(headers):
                val = r[idx] if idx < len(r) and r[idx] is not None else ''
                s_val = str(val).strip()
                if s_val.lower() == 'nan':
                    s_val = ''
                row_dict[col_name] = s_val
            dict_rows.append(row_dict)
        return headers, dict_rows

    @staticmethod
    def parse_uploaded_file(module_id: str, file_path_or_bytes: Any) -> Tuple[bool, List[Dict[str, Any]], str]:
        """两趟严格对齐解析用户上传的 WPS / Excel / CSV 表格"""
        try:
            headers, dict_rows = TemplateManager._read_rows_from_input(file_path_or_bytes)
            if not headers:
                return False, [], "文件中未检测到表头信息。"

            def normalize_str(s):
                return re.sub(r'[\s\(\)\[\]（）_\-]', '', str(s)).lower()

            def find_col_val(row: Dict[str, str], aliases: List[str]) -> str:
                # 第一趟扫描：严格全等匹配
                for alias in aliases:
                    for col in headers:
                        if col.strip().lower() == alias.strip().lower():
                            v = str(row.get(col, '')).strip()
                            if v:
                                return v
                # 第二趟扫描：去标点归一化严格匹配
                for alias in aliases:
                    norm_alias = normalize_str(alias)
                    for col in headers:
                        if normalize_str(col) == norm_alias:
                            v = str(row.get(col, '')).strip()
                            if v:
                                return v
                return ""

            records = []
            if module_id == "android":
                for row in dict_rows:
                    rec = {
                        '品牌编号': find_col_val(row, ['品牌编号', '品牌ID', '品牌代码']),
                        '品牌名称': find_col_val(row, ['品牌名称', '品牌', '厂商']),
                        '机型编号': find_col_val(row, ['机型编号', '机型ID', '型号', '型号ID']),
                        '内存': find_col_val(row, ['内存', '内存配置', '运行内存']),
                        '购买渠道': find_col_val(row, ['购买渠道', '渠道', '版本']),
                        '网络制式': find_col_val(row, ['网络制式', '网络', '制式']),
                        '颜色': find_col_val(row, ['颜色', '手机颜色', '机型颜色']),
                        '小型号': find_col_val(row, ['小型号', '苹果小型号', '副名称']),
                        '上市时间': find_col_val(row, ['上市时间', '发布时间', '时间', '上市日期']),
                        '开机价': find_col_val(row, ['开机价', '开机价格']),
                        '关机价': find_col_val(row, ['关机价', '关机价格']),
                        '开机最低价': find_col_val(row, ['开机最低价', '最低价', '开机最低'])
                    }
                    if rec.get('品牌编号') or rec.get('机型编号') or rec.get('品牌名称'):
                        records.append(rec)

            elif module_id == "laptop":
                for row in dict_rows:
                    rec = {
                        '品牌ID': find_col_val(row, ['品牌ID', '品牌编号']),
                        '品牌名称': find_col_val(row, ['品牌名称', '品牌', '厂商']),
                        '别名': find_col_val(row, ['别名', '副名称', '机型别名']),
                        '型号': find_col_val(row, ['型号', '型号ID', '机型编号']),
                        '固态硬盘': find_col_val(row, ['固态硬盘', '固态', 'SSD']),
                        '机械硬盘': find_col_val(row, ['机械硬盘', '机械', 'HDD']),
                        '显卡型号': find_col_val(row, ['显卡型号', '显卡', 'GPU']),
                        '处理器': find_col_val(row, ['处理器', 'CPU', '处理器型号']),
                        '开机价': find_col_val(row, ['开机价', '开机价格']),
                        '关机价': find_col_val(row, ['关机价', '关机价格']),
                        '开机最低价': find_col_val(row, ['开机最低价', '最低价', '开机最低']),
                        '发布时间': find_col_val(row, ['发布时间', '上市时间', '时间'])
                    }
                    if rec.get('品牌ID') or rec.get('型号') or rec.get('品牌名称'):
                        records.append(rec)

            elif module_id == "watch":
                for row in dict_rows:
                    rec = {
                        '机型编号': find_col_val(row, ['机型编号', '型号ID']),
                        '品牌ID': find_col_val(row, ['品牌ID', '品牌编号']),
                        '品牌名称': find_col_val(row, ['品牌名称', '品牌', '厂商']),
                        '型号': find_col_val(row, ['型号', '别名', '机型名称', '副名称']),
                        '购买渠道': find_col_val(row, ['购买渠道', '渠道', '版本']),
                        '手机颜色': find_col_val(row, ['手机颜色', '颜色', '机型颜色']),
                        '表壳材质': find_col_val(row, ['表壳材质', '材质']),
                        '表壳尺寸': find_col_val(row, ['表壳尺寸', '尺寸']),
                        '表带': find_col_val(row, ['表带', '表带类型']),
                        '连接': find_col_val(row, ['连接', '连接方式']),
                        '开机价': find_col_val(row, ['开机价', '开机价格']),
                        '关机价': find_col_val(row, ['关机价', '关机价格']),
                        '开机最低价': find_col_val(row, ['开机最低价', '最低价', '开机最低']),
                        '发布时间': find_col_val(row, ['发布时间', '上市时间', '时间'])
                    }
                    if rec.get('机型编号') or rec.get('品牌ID') or rec.get('品牌名称'):
                        records.append(rec)

            if not records:
                return False, [], "文件中未检测到有效机型数据，请核对表头。"

            return True, records, f"成功精准对齐并导入 {len(records)} 条机型配置数据！"
        except Exception as e:
            return False, [], f"解析表格失败: {str(e)}"
