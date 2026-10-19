# -*- coding: utf-8 -*-
"""
3C 机型字典库管理引擎 (Vercel Serverless 云端双保险版：内置全量字典 + Excel 动态增量读取)
"""
import os
import openpyxl
from typing import List, Dict, Optional
from config.builtin_dict import BUILTIN_BRANDS, BUILTIN_MOBILE, BUILTIN_LAPTOP, BUILTIN_WATCH

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class DictManager:
    def __init__(self):
        self.dict_file_paths = [
            os.path.join(BASE_DIR, "3C机型字典库.xlsx"),
            os.path.join(BASE_DIR, "data", "3C机型字典库.xlsx"),
            os.path.join(os.getcwd(), "3C机型字典库.xlsx"),
            os.path.join(os.getcwd(), "data", "3C机型字典库.xlsx"),
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "3C机型字典库.xlsx"),
            "/var/task/3C机型字典库.xlsx",
            "/var/task/data/3C机型字典库.xlsx",
        ]

        self.brands: List[Dict[str, str]] = []
        self.brand_name_to_id: Dict[str, str] = {}
        self.brand_id_to_name: Dict[str, str] = {}

        # 手机&平板 参数字典
        self.mobile_colors: List[str] = []
        self.mobile_networks: List[str] = []
        self.mobile_channels: List[str] = []
        self.mobile_rams: List[str] = []
        self.mobile_roms: List[str] = []

        # 笔记本 参数字典
        self.laptop_ssds: List[str] = []
        self.laptop_hdds: List[str] = []
        self.laptop_cpus: List[str] = []
        self.laptop_gpus: List[str] = []

        # 手表 参数字典
        self.watch_materials: List[str] = []
        self.watch_sizes: List[str] = []
        self.watch_connections: List[str] = []
        self.watch_straps: List[str] = []
        self.watch_networks: List[str] = []

        self.load_dictionary()

    def find_dict_path(self) -> Optional[str]:
        for p in self.dict_file_paths:
            try:
                if os.path.exists(p):
                    return p
            except Exception:
                continue
        return None

    def _add_brand(self, bid: str, bname: str):
        bid = str(bid or "").strip()
        if bid.endswith(".0"):
            bid = bid[:-2]
        bname = str(bname or "").strip()
        if not bid or not bname or bname.lower() == "nan":
            return
        if bname.lower() not in self.brand_name_to_id:
            self.brands.append({"id": bid, "name": bname})
            self.brand_name_to_id[bname.lower()] = bid
            self.brand_name_to_id[bname] = bid
            self.brand_id_to_name[bid] = bname

    def _load_builtin_fallback(self):
        """第一重保障：先加载内置全量字典（防止 Vercel Serverless 未打包 .xlsx 文件或路径偏移）"""
        for b in BUILTIN_BRANDS:
            self._add_brand(b.get("id", ""), b.get("name", ""))

        self.mobile_colors.extend(BUILTIN_MOBILE.get("机型颜色", []))
        self.mobile_networks.extend(BUILTIN_MOBILE.get("网络制式", []))
        self.mobile_channels.extend(BUILTIN_MOBILE.get("购买渠道", []))
        self.mobile_rams.extend(BUILTIN_MOBILE.get("RAM", []))
        self.mobile_roms.extend(BUILTIN_MOBILE.get("ROM", []))

        self.laptop_ssds.extend(BUILTIN_LAPTOP.get("固态硬盘", []))
        self.laptop_hdds.extend(BUILTIN_LAPTOP.get("机械硬盘", []))
        self.laptop_cpus.extend(BUILTIN_LAPTOP.get("处理器", []))
        self.laptop_gpus.extend(BUILTIN_LAPTOP.get("显卡", []))

        self.watch_materials.extend(BUILTIN_WATCH.get("表壳材质", []))
        self.watch_sizes.extend(BUILTIN_WATCH.get("表壳尺寸", []))
        self.watch_connections.extend(BUILTIN_WATCH.get("连接", []))
        self.watch_networks.extend(BUILTIN_WATCH.get("网络制式", []))

    def load_dictionary(self):
        """双保险加载：先加载内置完整字典，若存在《3C机型字典库.xlsx》则用 openpyxl 增量合并"""
        self._load_builtin_fallback()

        dict_path = self.find_dict_path()
        if dict_path:
            try:
                wb = openpyxl.load_workbook(dict_path, read_only=True, data_only=True)

                if "品牌id" in wb.sheetnames:
                    ws = wb["品牌id"]
                    rows = list(ws.iter_rows(values_only=True))
                    if rows:
                        headers = [str(h or "").strip() for h in rows[0]]
                        id_idx = headers.index("ID") if "ID" in headers else 0
                        name_idx = headers.index("品牌名称") if "品牌名称" in headers else 1
                        for r in rows[1:]:
                            if r and len(r) > max(id_idx, name_idx):
                                self._add_brand(r[id_idx], r[name_idx])

                def read_sheet_col(sheet_name: str, col_name: str) -> List[str]:
                    if sheet_name not in wb.sheetnames:
                        return []
                    ws = wb[sheet_name]
                    rows = list(ws.iter_rows(values_only=True))
                    if not rows:
                        return []
                    headers = [str(h or "").strip() for h in rows[0]]
                    if col_name not in headers:
                        return []
                    idx = headers.index(col_name)
                    vals = []
                    for r in rows[1:]:
                        if r and idx < len(r) and r[idx] is not None:
                            v = str(r[idx]).strip()
                            if v and v.lower() != "nan":
                                vals.append(v)
                    return vals

                phone_sheet = "手机&平板" if "手机&平板" in wb.sheetnames else ("手机" if "手机" in wb.sheetnames else "")
                if phone_sheet:
                    self.mobile_colors.extend(read_sheet_col(phone_sheet, "机型颜色"))
                    self.mobile_networks.extend(read_sheet_col(phone_sheet, "网络制式"))
                    self.mobile_channels.extend(read_sheet_col(phone_sheet, "购买渠道"))
                    self.mobile_rams.extend(read_sheet_col(phone_sheet, "RAM"))
                    self.mobile_roms.extend(read_sheet_col(phone_sheet, "ROM"))

                if "笔记本" in wb.sheetnames:
                    self.laptop_ssds.extend(read_sheet_col("笔记本", "固态硬盘"))
                    self.laptop_hdds.extend(read_sheet_col("笔记本", "机械硬盘"))
                    self.laptop_cpus.extend(read_sheet_col("笔记本", "处理器"))
                    self.laptop_gpus.extend(read_sheet_col("笔记本", "显卡"))

                if "手表" in wb.sheetnames:
                    self.watch_materials.extend(read_sheet_col("手表", "表壳材质"))
                    self.watch_sizes.extend(read_sheet_col("手表", "表壳尺寸"))
                    self.watch_connections.extend(read_sheet_col("手表", "连接"))
                    self.watch_networks.extend(read_sheet_col("手表", "网络制式"))

                wb.close()
            except Exception as e:
                print(f"读取 Excel 字典提示 (已自动使用内置全量字典兜底): {e}")

        # 默认手表表带等兜底
        if not self.watch_straps:
            self.watch_straps = ['TPU表带', '硅胶表带', '氟橡胶表带', '皮质表带', '真皮表带', '米兰尼斯表带', '不锈钢表带', '编织表带', '无表带', '有表带']
        if not self.watch_materials:
            self.watch_materials = ['铝金属表壳', '不锈钢表壳', '钛金属表壳', '精密陶瓷', '纤维高强度聚合物', '高分子材料', '陶瓷表壳', '纳米微晶陶瓷', '精钢表壳', '316L不锈钢材质']
        if not self.watch_sizes:
            self.watch_sizes = ['41毫米', '45毫米', '40毫米', '44毫米', '38毫米', '42毫米', '46毫米', '49毫米', '1.85寸', '1.43英寸', '1.91英寸', '2.07英寸']
        if not self.watch_connections:
            self.watch_connections = ['GPS', 'GPS+蜂窝网络', '蓝牙', 'WIFI', 'eSIM独立通话']

        self.mobile_colors = list(dict.fromkeys(self.mobile_colors))
        self.mobile_networks = list(dict.fromkeys(self.mobile_networks))
        self.mobile_channels = list(dict.fromkeys(self.mobile_channels))
        self.mobile_rams = list(dict.fromkeys(self.mobile_rams))
        self.mobile_roms = list(dict.fromkeys(self.mobile_roms))
        self.laptop_ssds = list(dict.fromkeys(self.laptop_ssds))
        self.laptop_hdds = list(dict.fromkeys(self.laptop_hdds))
        self.laptop_cpus = list(dict.fromkeys(self.laptop_cpus))
        self.laptop_gpus = list(dict.fromkeys(self.laptop_gpus))
        self.watch_materials = list(dict.fromkeys(self.watch_materials))
        self.watch_sizes = list(dict.fromkeys(self.watch_sizes))
        self.watch_connections = list(dict.fromkeys(self.watch_connections))
        self.watch_straps = list(dict.fromkeys(self.watch_straps))
        self.watch_networks = list(dict.fromkeys(self.watch_networks))


dict_manager = DictManager()
